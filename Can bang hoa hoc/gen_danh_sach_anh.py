# Danh sách ảnh của chuyên đề: vị trí dùng, trạng thái, nguồn, việc cần làm.
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROWS = [
    ['Hình 1.1. Tốc độ phản ứng thuận, nghịch theo thời gian', 'Buổi 1 / PX2', '✓ Đã làm nét', 'SGK Hóa học 11, KNTT', ''],
    ['Đồ thị nồng độ đi tới cân bằng từ hai phía', 'Buổi 1 / bài giảng PX1–PX2', '✓ Vẽ bằng matplotlib theo số liệu PX1 (Kc = 64)', 'P20 STUDY vẽ; tham khảo SCH 4U Strand 3, Unit 5 Chemical Equilibrium', 'Chỉ dùng trong bài giảng, không đặt trong PX1 (lộ đáp án ý 4)'],
    ['Thạch nhũ và măng đá', 'Buổi 1 / PX1', 'Chờ Gemini dịch nhãn', 'Voarintsoa, N. R. G., ResearchGate, Figure 7 (Việt hóa)', 'Cắt tay phần chú thích tiếng Anh dưới hình trước khi gửi Gemini'],
    ['Vòi nước (cân bằng động) và bập bênh (cân bằng tĩnh)', 'Buổi 1 / PX2', 'Chờ Gemini vẽ lại theo mẫu', 'Vẽ lại theo mẫu: "Equilibria (IB or A level chemistry)", Lmbroomfield, tes.com; bập bênh: world-builders.org', 'Tài liệu gốc trả phí, chỉ dùng làm mẫu bố cục'],
    ['Hình 1.2. Nhiệt độ và cân bằng 2NO₂ ⇌ N₂O₄', 'Buổi 2 / PX9', '✓ Đã làm nét', 'SGK Hóa học 11, KNTT', 'Nền tím xám, lệch tông với 1.3, 1.4'],
    ['Hình 1.3. Nhiệt độ và cân bằng thủy phân CH₃COONa', 'Buổi 2 / PX9', '✓ Đã làm nét', 'SGK Hóa học 11, KNTT', 'Ống (2) gần cùng màu ống (1); nói rõ khi giảng'],
    ['Hình 1.4. Nồng độ và cân bằng thủy phân CH₃COONa', 'Buổi 2 / PX7', '✓ Đã làm nét', 'SGK Hóa học 11, KNTT', ''],
    ['Hình 1.5. Áp suất và cân bằng NO₂ trong xi-lanh', 'Buổi 2 / PX8', '✓ Đã làm nét', 'SGK Hóa học 11, KNTT', 'Ghi chú dưới hình: sau khi nén, màu nhạt bớt nhưng vẫn đậm hơn lúc chưa nén'],
    ['Sơ đồ dây chuyền tổng hợp NH₃ (vòng tuần hoàn)', 'Buổi 3 / Đề 2', 'Chờ Gemini vẽ lại theo mẫu', 'Vẽ lại theo mẫu: "Equilibria (IB or A level chemistry)", Lmbroomfield, tes.com', 'Sửa ΔH về −91,8 kJ theo SGK'],
]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
s = doc.sections[0]
s.page_width, s.page_height = Cm(29.7), Cm(21)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('CHUYÊN ĐỀ CÂN BẰNG HÓA HỌC: DANH SÁCH ẢNH VÀ NGUỒN'); r.bold = True; r.font.size = Pt(14)
t = doc.add_table(rows=1, cols=5); t.style = 'Table Grid'
for c, h in zip(t.rows[0].cells, ['Ảnh', 'Dùng ở', 'Trạng thái', 'Nguồn', 'Ghi chú']):
    c.text = ''; c.paragraphs[0].add_run(h).bold = True
for row in ROWS:
    cells = t.add_row().cells
    for c, v in zip(cells, row): c.text = v
doc.add_paragraph()
p = doc.add_paragraph(); r = p.add_run('Cách ghi nguồn dưới hình: '); r.bold = True
p.add_run('Nguồn: SGK Hóa học 11, KNTT. Ảnh nước ngoài ghi tác giả hoặc trang gốc, kèm "Việt hóa" hoặc "vẽ lại theo mẫu".').italic = True
doc.save('Danh sách ảnh và nguồn.docx'); print('ok')
