# Sinh docx GV Đề 1 (ngữ liệu + 6 câu + đáp án, giải thích ngay dưới từng câu), format A13/A14.
import os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
RED, ORANGE, GRAY, GREEN = RGBColor(0xEE, 0, 0), RGBColor(0xFF, 0x80, 0), RGBColor(0x80, 0x80, 0x80), RGBColor(0, 0x80, 0)
C, L, J, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT
SUB = dict(zip('₀₁₂₃₄₅₆₇₈₉', '0123456789'))
SUP = dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺', '0123456789-+'))
TOKEN = re.compile('([₀-₉]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺]+)')

TITLE = 'KHI LE CHATELIER "ĐOÁN SAI"'
IMG = {
    '1': os.path.join(HERE, 'image', 'de1_hinh2.png'),  # khí trơ A1/A2
    '2': os.path.join(HERE, 'image', 'de1_hinh1.png'),  # FeSCN pha loãng
}
PASSAGE = [
    ('p', '[1] Nguyên lí Le Chatelier thường được học sinh nhớ dưới dạng hai câu ngắn: "tăng áp suất thì cân bằng dịch chuyển về phía ít phân tử khí" và "thêm chất nào thì hệ tiêu thụ bớt chất đó". Hai câu này chỉ là cách nói gọn của một quy tắc định lượng. Gọi Q là tỉ số tính bằng đúng công thức của Kc nhưng thay vào nồng độ ngay sau tác động. Nếu Q < Kc, cân bằng chuyển dịch theo chiều thuận; nếu Q > Kc, cân bằng chuyển dịch theo chiều nghịch; nếu Q = Kc, cân bằng không dịch chuyển. Kc chỉ thay đổi khi nhiệt độ thay đổi, và nồng độ chất rắn không có mặt trong biểu thức Kc.'),
    ('p', '[2] Tình huống A. Một hỗn hợp khí đang ở trạng thái cân bằng N₂O₄(g) ⇌ 2NO₂(g) ở nhiệt độ T, với Kc = [NO₂]²/[N₂O₄] = 0,036, [N₂O₄] = 0,100 M và [NO₂] = 0,060 M. Người ta thêm khí argon (Ar), một khí không phản ứng với N₂O₄ và NO₂, theo hai cách. Cách A1: hỗn hợp chứa trong bình cứng, thể tích không đổi; sau khi thêm Ar, áp suất tổng trong bình tăng. Cách A2: hỗn hợp chứa trong xi-lanh có pít-tông chuyển động tự do nên áp suất tổng giữ không đổi; lượng Ar thêm vào làm thể tích khí tăng gấp đôi.'),
    ('fig', '1', 'Thêm khí trơ Ar vào bình cứng (A1) và vào xi-lanh có pít-tông tự do (A2). Nguồn: Shutterstock, chỉnh sửa bởi P20 STUDY.'),
    ('p', '[3] Tình huống B. Một dung dịch chứa cân bằng Fe³⁺(aq) + SCN⁻(aq) ⇌ FeSCN²⁺(aq); màu đỏ máu của dung dịch do ion FeSCN²⁺ tạo ra. Ở 25 °C, Kc = [FeSCN²⁺]/([Fe³⁺][SCN⁻]) = 100. Ban đầu [Fe³⁺] = [SCN⁻] = 0,020 M và [FeSCN²⁺] = 0,040 M. Người ta thêm nước cất để thể tích dung dịch tăng gấp đôi, rồi dùng máy đo màu theo dõi nồng độ FeSCN²⁺ đến khi số đo ổn định.'),
    ('fig', '2', 'Dung dịch chứa FeSCN²⁺ trước và sau khi pha loãng gấp đôi. Nguồn: P20 STUDY.'),
    ('p', '[4] Tình huống C. Một bình kín ở 25 °C chứa dư NH₄HS rắn, có cân bằng NH₄HS(s) ⇌ NH₃(g) + H₂S(g) với Kc = [NH₃][H₂S] = 1,44·10⁻⁴; lúc cân bằng [NH₃] = [H₂S] = 0,012 M. Người ta cho thêm 10 g NH₄HS rắn vào bình, giữ nguyên nhiệt độ và thể tích.'),
    ('tab', '1', 'Nồng độ (M) các chất ở ba thời điểm.', [
        ['Tình huống', 'Chất', 'Cân bằng ban đầu', 'Ngay sau tác động', 'Cân bằng mới'],
        ['A1', 'N₂O₄', '0,100', '0,100', ''], ['A1', 'NO₂', '0,060', '0,060', ''],
        ['A2', 'N₂O₄', '0,100', '0,050', '0,0449'], ['A2', 'NO₂', '0,060', '0,030', '0,0402'],
        ['B', 'Fe³⁺', '0,020', '0,010', '0,0130'], ['B', 'SCN⁻', '0,020', '0,010', '0,0130'],
        ['B', 'FeSCN²⁺', '0,040', '0,020', '0,0170'],
        ['C', 'NH₃', '0,012', '0,012', ''], ['C', 'H₂S', '0,012', '0,012', '']]),
    ('tab', '2', 'Dự đoán của một nhóm học sinh khi chỉ dựa vào phát biểu dạng lời của nguyên lí Le Chatelier.', [
        ['Tác động', 'Dự đoán của nhóm học sinh'],
        ['Thêm Ar vào bình cứng (A1)', 'Áp suất tổng tăng nên cân bằng chuyển dịch về phía ít phân tử khí, tức chiều nghịch'],
        ['Thêm Ar ở áp suất không đổi (A2)', 'Áp suất tổng không đổi nên cân bằng không dịch chuyển'],
        ['Pha loãng dung dịch (B)', 'Mọi nồng độ cùng giảm một nửa nên cân bằng không dịch chuyển'],
        ['Thêm NH₄HS rắn (C)', 'Thêm chất đầu nên cân bằng chuyển dịch theo chiều thuận'],
        ['Rút bớt NO₂ khỏi bình cứng ở tình huống A (trước khi thêm Ar)', 'Mất bớt sản phẩm nên cân bằng chuyển dịch theo chiều thuận để tạo lại NO₂']]),
]

# Mỗi câu: meta, stem, kiểu + nội dung, đáp án, giải thích (list dòng)
Q = [
    dict(meta='Trắc nghiệm 4 lựa chọn | Độ khó 1,0/5 | Mức tư duy 1 | Lỗi khai thác: L5, L3', kind='mcq',
         stem='Theo đoạn [1], phát biểu nào sau đây về việc so sánh Q với Kc là đúng?',
         opts=['Cân bằng chuyển dịch nhằm mục đích làm thay đổi giá trị hằng số cân bằng Kc của hệ.',
               'Cân bằng chuyển dịch theo chiều thuận khi giá trị Q tính được lớn hơn Kc.',
               'Cân bằng không chuyển dịch nếu Q tính được ngay sau tác động bằng Kc.',
               'Cân bằng chuyển dịch theo chiều thuận khi hệ được thêm một chất đầu bất kỳ.'],
         ans='C', exp=['Đoạn [1] nêu: nếu Q = Kc, cân bằng không dịch chuyển. Vậy C đúng.',
                       'A sai: Kc chỉ thay đổi khi nhiệt độ thay đổi; cân bằng chuyển dịch để Q tiến về Kc, không làm Kc thay đổi (L5).',
                       'B sai: Q > Kc thì cân bằng chuyển dịch theo chiều nghịch.',
                       'D sai: nếu chất đầu là chất rắn thì Q không đổi nên cân bằng không dịch chuyển (L3).']),
    dict(meta='Đúng–Sai | Độ khó 2,5/5 | Mức tư duy 2 | Lỗi khai thác: L1, L4', kind='tf',
         stem='Các nhận định sau đúng hay sai?',
         items=['Trong tình huống A1, Q ngay sau tác động bằng 0,036.',
                'Thêm Ar vào bình cứng (A1) làm cân bằng chuyển dịch theo chiều nghịch.',
                'Trong tình huống A2, Q ngay sau tác động lớn hơn Kc.',
                'Ở tình huống A2, cân bằng bị phá vỡ vì thể tích tăng làm nồng độ N₂O₄ và NO₂ giảm, dù áp suất tổng không đổi.'],
         ans='(a) Đúng; (b) Sai; (c) Sai; (d) Đúng',
         exp=['(a) Đúng: từ Bảng 1, ngay sau tác động [NO₂] = 0,060 M, [N₂O₄] = 0,100 M nên Q = 0,060²/0,100 = 0,036.',
              '(b) Sai: Q = Kc nên cân bằng không dịch chuyển. Áp suất tổng tăng nhưng nồng độ các chất phản ứng giữ nguyên (L1).',
              '(c) Sai: ở A2, Q = 0,030²/0,050 = 0,018, nhỏ hơn Kc = 0,036.',
              '(d) Đúng: thể tích tăng gấp đôi làm nồng độ giảm còn 0,050 M và 0,030 M, Q lệch khỏi Kc nên cân bằng bị phá vỡ. Lỗi L4 là tưởng áp suất tổng không đổi thì nồng độ không đổi.']),
    dict(meta='Trắc nghiệm 4 lựa chọn | Độ khó 2,0/5 | Mức tư duy 2 | Lỗi khai thác: L2', kind='mcq',
         stem='Dựa vào số liệu tình huống B, Q ngay sau khi pha loãng gấp đôi và chiều chuyển dịch của cân bằng là',
         opts=['Q = 200; cân bằng chuyển dịch theo chiều nghịch.', 'Q = 50; cân bằng chuyển dịch theo chiều thuận.',
               'Q = 100; cân bằng vẫn giữ nguyên, không chuyển dịch.', 'Q = 200; cân bằng chuyển dịch theo chiều thuận.'],
         ans='A', exp=['Ngay sau pha loãng: [FeSCN²⁺] = 0,020 M; [Fe³⁺] = [SCN⁻] = 0,010 M.',
                       'Q = [FeSCN²⁺]/([Fe³⁺][SCN⁻]) = 0,020/(0,010 × 0,010) = 200.',
                       'Q = 200 > Kc = 100 nên cân bằng chuyển dịch theo chiều nghịch. Vậy A đúng.',
                       'B sai: đặt nhầm tử số và mẫu số.',
                       'C sai: cho rằng nồng độ giảm đều thì Q giữ nguyên (L2).',
                       'D sai: tính đúng Q nhưng so sánh nhầm chiều chuyển dịch.']),
    dict(meta='Đúng–Sai | Độ khó 2,5/5 | Mức tư duy 3 | Giả định tình huống | Lỗi khai thác: L5', kind='tf',
         stem='Ở tình huống C, giả sử thay vì thêm NH₄HS rắn, người ta bơm thêm khí H₂S vào bình (giữ nguyên nhiệt độ và thể tích). Các nhận định sau đúng hay sai?',
         items=['Q ngay sau khi bơm H₂S lớn hơn Kc.', 'Cân bằng chuyển dịch theo chiều nghịch, làm tăng lượng NH₄HS rắn.',
                'Ở trạng thái cân bằng mới, nồng độ NH₃ lớn hơn 0,012 M.', 'Bơm thêm H₂S làm hằng số cân bằng Kc tăng lên.'],
         ans='(a) Đúng; (b) Đúng; (c) Sai; (d) Sai',
         exp=['(a) Đúng: Q = [NH₃][H₂S]; [H₂S] tăng còn [NH₃] giữ nguyên nên Q > Kc.',
              '(b) Đúng: Q > Kc nên cân bằng chuyển dịch theo chiều nghịch, NH₃ và H₂S kết hợp thành NH₄HS rắn.',
              '(c) Sai: chiều nghịch tiêu thụ NH₃ nên ở cân bằng mới [NH₃] nhỏ hơn 0,012 M.',
              '(d) Sai: Kc chỉ phụ thuộc nhiệt độ, không đổi khi thêm chất (L5).']),
    dict(meta='Ghép đôi | Độ khó 4,0/5 | Mức tư duy 4 | Chẩn đoán & Khắc phục | Lỗi khai thác: L1, L2, L3, L4', kind='match',
         stem='Mỗi dự đoán sai trong Bảng 2 xuất phát từ một lỗi lập luận. Ghép mỗi ý ở Cột A với một ý ở Cột B.',
         head=['Cột A', 'Cột B'],
         A=['1. Dự đoán cho tình huống A1', '2. Dự đoán cho tình huống A2', '3. Dự đoán cho tình huống B', '4. Dự đoán cho tình huống C'],
         B=['a. Bỏ qua sự khác nhau về số mũ và số lượng nồng độ ở tử số, mẫu số của biểu thức Q.',
            'b. Áp dụng máy móc quy tắc "thêm chất nào hệ tiêu thụ chất đó" cho chất rắn.',
            'c. Đồng nhất áp suất tổng tăng với nồng độ các chất phản ứng tăng.',
            'd. Đồng nhất áp suất tổng không đổi với nồng độ các chất không đổi.'],
         ans='1-c; 2-d; 3-a; 4-b',
         exp=['1-c: khí trơ trong bình cứng làm áp suất tổng tăng nhưng nồng độ N₂O₄, NO₂ không đổi (L1).',
              '2-d: áp suất tổng không đổi nhưng thể tích tăng làm loãng các khí, Q thay đổi (L4).',
              '3-a: tử số bậc một, mẫu số là tích hai nồng độ nên Q không giữ nguyên khi pha loãng (L2).',
              '4-b: chất rắn không có mặt trong Q nên thêm chất rắn không làm cân bằng dịch chuyển (L3).']),
    dict(meta='Kéo thả | Độ khó 3,5/5 | Mức tư duy 3 | Phân tích dữ liệu chéo | Lỗi khai thác: L2', kind='drag',
         stem='Kéo các cụm từ trong kho vào chỗ trống cho phù hợp:',
         bank='một nửa || gấp đôi || lớn hơn || nhỏ hơn || tiếp tục giảm || tăng nhẹ trở lại || thuận || nghịch',
         body='Trong tình huống B, ngay sau khi pha loãng gấp đôi, nồng độ FeSCN²⁺ còn (1) ______ giá trị ban đầu. Nếu pha loãng không làm cân bằng dịch chuyển, số đo của máy đo màu sẽ giữ nguyên giá trị này. Số đo thực tế khi ổn định là 0,0170 M, tức nồng độ FeSCN²⁺ đã (2) ______. Số liệu này cho thấy cân bằng đã chuyển dịch theo chiều (3) ______, phù hợp với việc Q ngay sau tác động (4) ______ Kc.',
         ans='(1) một nửa; (2) tiếp tục giảm; (3) nghịch; (4) lớn hơn', exp=[]),
]


class W:
    def __init__(self):
        self.d = Document()
        st = self.d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.paragraph_format.space_after = Pt(3)
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)

    def runs(self, p, text, fmt='', size=None):
        for part in TOKEN.split(text):
            if not part:
                continue
            sub, sup = part[0] in SUB, part[0] in SUP
            r = p.add_run(''.join((SUB if sub else SUP)[c] for c in part) if (sub or sup) else part)
            r.font.subscript, r.font.superscript = sub, sup
            r.bold, r.italic = 'b' in fmt.split(), 'i' in fmt.split()
            for k, col in (('red', RED), ('org', ORANGE), ('gray', GRAY), ('grn', GREEN)):
                if k in fmt.split():
                    r.font.color.rgb = col
            if size:
                r.font.size = Pt(size)

    def para(self, parts, align=J, size=None, indent=None, cell=None):
        p = cell.add_paragraph() if cell is not None else self.d.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        for t, f in parts:
            self.runs(p, t, f, size)
        return p

    def cellp(self, cell, text, fmt='', align=L):
        p = cell.paragraphs[0]; p.alignment = align; self.runs(p, text, fmt)

    def no_border_row(self, row):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
            for side in ('top', 'left', 'right'):
                e = OxmlElement('w:' + side); e.set(qn('w:val'), 'nil'); b.append(e)
            tcPr.append(b)

    def table(self, rows, widths=None, head_center=True):
        t = self.d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                self.cellp(t.cell(i, j), v, 'b' if i == 0 else '', C if (i == 0 and head_center) or len(rows[0]) > 2 else L)
        if widths:
            # LibreOffice/Word đọc độ rộng từ tblGrid, nên phải đặt cả cột lẫn từng ô và tắt autofit
            t.autofit = False
            for col, wd in zip(t.columns, widths):
                col.width = Cm(wd)
            for row in t.rows:
                for c, wd in zip(row.cells, widths):
                    c.width = Cm(wd)
        return t


def build():
    w = W()
    w.para([('CHỦ ĐỀ 1. ', 'b'), (TITLE, 'b red')], C, 13)
    for b in PASSAGE:
        if b[0] == 'p':
            num, txt = b[1].split(' ', 1)
            w.para([(num + ' ', 'b'), (txt, '')])
        elif b[0] == 'fig':
            img_path = IMG.get(b[1])
            if img_path and os.path.isfile(img_path):
                p = w.d.add_paragraph()
                p.alignment = C
                r = p.add_run()
                r.add_picture(img_path, width=Cm(13))
            else:
                w.para([('[Chèn Hình %s]' % b[1], 'i gray')], C)
            w.para([('Hình %s' % b[1], 'b i'), (': ' + b[2], 'i')], C)
        else:
            w.para([('Bảng %s' % b[1], 'b i'), (': ' + b[2], 'i')], C)
            w.table(b[3], widths=[6.5, 10] if len(b[3][0]) == 2 else None)
            w.d.add_paragraph()
    w.para([('Nguồn: P20 STUDY sưu tầm và tổng hợp', 'b')], R)
    w.para([('Đọc ngữ liệu trên và trả lời các câu hỏi sau đây:', 'i')], L)

    for n, q in enumerate(Q, 1):
        w.para([('Câu %d. [P20 STUDY]' % n, 'b org')], L)
        w.para([(q['meta'], 'i gray')], L, 10)
        w.para([(q['stem'], '')])
        if q['kind'] == 'mcq':
            if q.get('short'):
                w.para([('   '.join('%s. %s' % ('ABCD'[i], o) for i, o in enumerate(q['opts'])), '')], L)
            else:
                for i, o in enumerate(q['opts']):
                    w.para([('%s. ' % 'ABCD'[i], 'b'), (o, '')], L)
        elif q['kind'] == 'tf':
            t = w.table([['', 'Đúng', 'Sai']] + [['(%s) %s' % ('abcd'[i], it), '', ''] for i, it in enumerate(q['items'])],
                        widths=[13, 2, 2])
            w.no_border_row(t.rows[0])
            for r in t.rows[1:]:
                r.cells[0].paragraphs[0].alignment = L
        elif q['kind'] == 'fill':
            w.para([(q['body'], '')])
        elif q['kind'] == 'match':
            rows = [q['head']] + [[a, b] for a, b in zip(q['A'], q['B'])]
            t = w.table(rows, widths=[8.5, 8.5])
            for r in t.rows[1:]:
                for c in r.cells:
                    c.paragraphs[0].alignment = L
        elif q['kind'] == 'drag':
            w.para([('Kho: ', 'b'), (q['bank'], '')], C)
            w.para([(q['body'], '')])
        w.para([('▸ Đáp án: ', 'b grn'), (q['ans'], 'grn')], L)
        if q['exp']:
            w.para([('▸ Giải thích:', 'b grn')], L)
            for line in q['exp']:
                w.para([(line, 'grn')], J, indent=0.6)

    w.para([('Bảng tổng hợp', 'b i'), (': Đề 1 sau khi lọc.', 'i')], C)
    rows = [['Câu', 'Dạng', 'Độ khó', 'Mức', 'Kỹ năng', 'Lỗi L', 'Đáp án']]
    for n, q in enumerate(Q, 1):
        parts = [x.strip() for x in q['meta'].split('|')]
        skill = parts[3] if len(parts) == 5 else ''
        rows.append([str(n), parts[0], parts[1].replace('Độ khó ', ''), parts[2].replace('Mức tư duy ', ''),
                     skill, parts[-1].replace('Lỗi khai thác: ', ''), q['ans']])
    w.table(rows)
    out = os.path.join(HERE, 'Buổi 3. [SOURCE] Đề 1 - Khi Le Chatelier đoán sai.docx')
    w.d.save(out); return out


if __name__ == '__main__':
    print(build())
