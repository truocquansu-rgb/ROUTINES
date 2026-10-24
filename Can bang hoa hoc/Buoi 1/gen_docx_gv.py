# Sinh bản GV [SOURCE] Buổi 1 từ buoi1_spec.py (format A13/A14).
import sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from buoi1_spec import TITLE, LESSON, PX, DEBATE

RED, ORANGE, GRAY, GREEN = (RGBColor(0xEE, 0, 0), RGBColor(0xFF, 0x80, 0),
                            RGBColor(0x80, 0x80, 0x80), RGBColor(0, 0x80, 0))
C, L, J, R = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
              WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT)

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
st.paragraph_format.space_after = Pt(4)
for s in doc.sections:
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)


def para(runs, align=L, size=None):
    p = doc.add_paragraph(); p.alignment = align
    for text, fmt in runs:
        r = p.add_run(text)
        r.bold = 'b' in fmt; r.italic = 'i' in fmt
        if 'red' in fmt: r.font.color.rgb = RED
        if 'org' in fmt: r.font.color.rgb = ORANGE
        if 'gray' in fmt: r.font.color.rgb = GRAY
        if 'grn' in fmt: r.font.color.rgb = GREEN
        if size: r.font.size = Pt(size)
    return p


def table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j); cell.text = ''
            p = cell.paragraphs[0]; p.alignment = C
            r = p.add_run(val); r.bold = i == 0
    doc.add_paragraph()


para([(TITLE, 'b')], C, 14)
para([(LESSON, 'b red')], C, 13)
para([('Bản GV [SOURCE]: bài tập phản xạ PX1–PX6, đáp án và giải thích ngay sau từng ý', 'i gray')], C, 10)

for px in PX:
    code, name = px['title'].split('. ', 1)
    para([(code + '. ', 'b'), (name, 'b red')], L, 13)
    para([('Ý cần giảng (GV):', 'b i gray')], L, 10)
    for t in px['teach']:
        para([('• ' + t, 'i gray')], L, 10)
    para([('Ngữ liệu lạ', 'b i')], L)
    for b in px['passage']:
        if isinstance(b, str):
            para([(b, '')], J)
        else:
            para([('Bảng ' + b['label'], 'b i'), (': ' + b['cap'], 'i')], C)
            table(b['rows'])
    para([('Câu hỏi:', 'b')], L)
    for k, q in enumerate(px['questions'], 1):
        para([('(%d) ' % k, 'b org'), (q['q'], '')], J)
        para([('▸ Đáp án: ', 'b grn'), (q['ans'], 'grn')], J)
        para([('▸ Giải thích: ', 'b grn'), (q['exp'], 'grn')], J)

para([('CÂU HỎI TRANH BIỆN CỦA BUỔI', 'b red')], C, 13)
for k, d in enumerate(DEBATE, 1):
    para([('%d. ' % k, 'b'), (d, '')], J)

para([('Bảng tổng hợp', 'b i'), (': Bài tập phản xạ Buổi 1.', 'i')], C)
table([['PX', 'Cơ chế', 'Số ý']] +
      [[p['title'].split('. ', 1)[0], p['title'].split('. ', 1)[1], str(len(p['questions']))] for p in PX])

para([('Nguồn: P20 STUDY sưu tầm và tổng hợp', 'b')], R)

out = os.path.join(HERE, 'Buổi 1. [SOURCE] Bài tập phản xạ PX1–PX6.docx')
doc.save(out); print(out)
