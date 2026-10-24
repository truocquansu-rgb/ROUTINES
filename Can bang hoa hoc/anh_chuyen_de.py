# Danh mục ảnh chuyên đề (một nguồn duy nhất): sinh zip ảnh đã duyệt và file danh sách ảnh.
import os, zipfile
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

KNTT = 'SGK Hóa học 11, KNTT'
# (buổi, PX, file, tên hình, nguồn, ghi chú); file None = chưa có
ANH = [
    (1, 'PX1', 'Buoi 1/image/hang_dong_phong_nha.png', 'Hang động có thạch nhũ (mở bài)', 'VOV, "New cave found in Vietnam\'s Phong Nha-Ke Bang reveals rare geological wonders"', ''),
    (1, 'PX1', 'Buoi 1/image/thach_nhu.png', 'Thạch nhũ và măng đá, phản ứng thuận nghịch trong tự nhiên', 'Voarintsoa, N. R. G., ResearchGate, Figure 7 (Việt hóa)', ''),
    (1, 'PX1', 'Buoi 1/image/do_thi_nong_do_hai_chieu.png', 'Đồ thị nồng độ đi tới cân bằng từ hai phía', 'P20 STUDY vẽ theo số liệu PX1; tham khảo SCH 4U Unit 5 Chemical Equilibrium', 'Chỉ dùng trong bài giảng, không đặt trong đề PX1'),
    (1, 'PX2', 'Buoi 1/image/hinh_1_1.png', 'Tốc độ phản ứng thuận và nghịch theo thời gian', KNTT, ''),
    (1, 'PX2', 'Buoi 1/image/voi_nuoc_bap_benh.png', 'Cân bằng động (vòi nước) và cân bằng tĩnh (bập bênh)', 'Vẽ lại theo mẫu: "Equilibria (IB or A level chemistry)", Lmbroomfield, tes.com', 'Chốt ý sau hình tốc độ'),
    (1, 'PX3', 'Buoi 1/image/lo_nung_voi.png', 'Lò nung vôi: quạt hút CO₂ để CaCO₃ phân hủy gần hết', 'Wikipedia, "Lime kiln" (Việt hóa)', ''),
    (1, 'PX5', 'Buoi 1/image/no_bong_H2.png', 'Nổ bóng H₂: Kc rất lớn nhưng cần tia lửa', 'ThoughtCo, "Hydrogen Balloon Explosion"', ''),
    (1, 'PX5', 'Buoi 1/image/bo_xuc_tac_o_to.png', 'Bộ xúc tác trên đường ống xả ô tô', 'DLR elib (Việt hóa)', ''),
    (1, 'PX6', 'Buoi 1/image/Q_va_K.png', 'So sánh Q với K quyết định chiều phản ứng', 'Chemistry Steps, "Reaction Quotient – Q" (Việt hóa)', ''),
    (2, 'PX7', 'Buoi 2/image/hinh_1_4.png', 'Nồng độ và cân bằng thủy phân CH₃COONa', KNTT, ''),
    (2, 'PX8', 'Buoi 2/image/hinh_1_5.png', 'Áp suất và cân bằng NO₂ trong xi-lanh', KNTT, 'Sau khi nén, màu nhạt bớt nhưng vẫn đậm hơn lúc chưa nén'),
    (2, 'PX9', 'Buoi 2/image/hinh_1_2.png', 'Nhiệt độ và cân bằng 2NO₂ ⇌ N₂O₄', KNTT, 'Dạy cặp với hình tiếp theo'),
    (2, 'PX9', 'Buoi 2/image/hinh_1_3.png', 'Nhiệt độ và cân bằng thủy phân CH₃COONa', KNTT, 'Ống (2) gần cùng màu ống (1), nói rõ khi giảng'),
    (3, 'Đề 2', None, 'Sơ đồ dây chuyền tổng hợp NH₃ (vòng tuần hoàn)', 'Vẽ lại theo mẫu: "Equilibria (IB or A level chemistry)", Lmbroomfield, tes.com', 'Sửa ΔH về −91,8 kJ theo SGK'),
]

SAFE = str.maketrans({'₂': '2', '₃': '3', '₄': '4', '⇌': '-', '/': '-', ':': ' -'})
idx, cho, so = ['THỨ TỰ HÌNH, VỊ TRÍ DÙNG VÀ NGUỒN', ''], [], {}
with zipfile.ZipFile('Cân bằng hóa học - Ảnh đã duyệt.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    for b, px, f, ten, nguon, note in ANH:
        so[b] = so.get(b, 0) + 1
        name = 'Buổi %d/Hình %d - %s - %s' % (b, so[b], px, ten.translate(SAFE))
        if f and os.path.exists(f):
            z.write(f, name + '.png')
            idx += [name, '   Nguồn: ' + nguon + ('\n   Ghi chú: ' + note if note else '')]
        else:
            cho.append(name)
    idx += ['', 'Còn chờ:'] + ['   ' + c for c in cho]
    z.writestr('Thứ tự hình và nguồn.txt', '\n'.join(idx))
print('\n'.join(idx))

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
s = doc.sections[0]; s.page_width, s.page_height = Cm(29.7), Cm(21)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('CHUYÊN ĐỀ CÂN BẰNG HÓA HỌC: DANH SÁCH ẢNH VÀ NGUỒN'); r.bold = True; r.font.size = Pt(14)
t = doc.add_table(rows=1, cols=6); t.style = 'Table Grid'
for c, h in zip(t.rows[0].cells, ['Buổi', 'Dùng ở', 'Ảnh', 'Trạng thái', 'Nguồn', 'Ghi chú']):
    c.text = ''; c.paragraphs[0].add_run(h).bold = True
for b, px, f, ten, nguon, note in ANH:
    ok = bool(f and os.path.exists(f))
    for c, v in zip(t.add_row().cells, [str(b), px, ten, '✓ Đã duyệt' if ok else 'Chờ', nguon, note]): c.text = v
doc.add_paragraph().add_run('Cách ghi nguồn dưới hình: "Nguồn: SGK Hóa học 11, KNTT". Ảnh nước ngoài ghi tác giả hoặc trang gốc, kèm "Việt hóa" hoặc "vẽ lại theo mẫu".').italic = True
doc.save('Danh sách ảnh và nguồn.docx')
