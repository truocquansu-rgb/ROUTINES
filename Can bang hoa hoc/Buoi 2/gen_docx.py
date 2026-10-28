# Sinh bài giảng Buổi 2 dạng docx: bản HS và bản GV [SOURCE].
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE); sys.path.insert(0, ROOT)
import buoi2_spec as S
from anh_chuyen_de import so_hinh

FIG = so_hinh()
RED, ORANGE, GRAY, GREEN, BLUE = (RGBColor(0xEE, 0, 0), RGBColor(0xE6, 0x51, 0x00), RGBColor(0x80, 0x80, 0x80),
                                  RGBColor(0, 0x80, 0), RGBColor(0x15, 0x65, 0xC0))
C, L, J, R = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.RIGHT)
SUB = dict(zip('₀₁₂₃₄₅₆₇₈₉', '0123456789'))
SUP = dict(zip('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺', '0123456789-+'))
TOKEN = re.compile('([₀-₉]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺]+)')


def fig_ref(text):
    return re.sub(r'\{H:([^}]+)\}', lambda m: str(FIG[m.group(1)][1]), text)


class Doc:
    def __init__(self, gv):
        self.gv = gv
        self.d = Document()
        st = self.d.styles['Normal']
        st.font.name = 'Times New Roman'; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        st.paragraph_format.space_after = Pt(4)
        for s in self.d.sections:
            s.page_width, s.page_height = Cm(21), Cm(29.7)
            s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2)

    def runs(self, p, text, fmt='', size=None):
        for part in TOKEN.split(fig_ref(text)):
            if not part:
                continue
            sub = part[0] in SUB; sup = part[0] in SUP
            r = p.add_run(''.join((SUB if sub else SUP).get(c, c) for c in part) if (sub or sup) else part)
            r.font.subscript = sub; r.font.superscript = sup
            r.bold = 'b' in fmt; r.italic = 'i' in fmt
            for key, col in (('red', RED), ('org', ORANGE), ('gray', GRAY), ('grn', GREEN), ('blue', BLUE)):
                if key in fmt.split():
                    r.font.color.rgb = col
            if size:
                r.font.size = Pt(size)

    def para(self, parts, align=J, size=None, indent=None):
        p = self.d.add_paragraph(); p.alignment = align
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        for text, fmt in parts:
            self.runs(p, text, fmt, size)
        return p

    def border_bottom(self, p, color='1565C0'):
        pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
        for k, v in (('val', 'single'), ('sz', '8'), ('space', '1'), ('color', color)):
            e.set(qn('w:' + k), v)
        b.append(e); pPr.append(b)

    def table(self, rows, label=None, cap=None):
        if label:
            self.para([('Bảng ' + label, 'b i'), (': ' + cap, 'i')], C)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j); cell.text = ''
                p = cell.paragraphs[0]; p.alignment = C
                self.runs(p, val, 'b' if i == 0 else '')
        self.d.add_paragraph()

    def figures(self, paths):
        t = self.d.add_table(rows=1, cols=len(paths)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        width = {1: 11.0, 2: 7.8, 3: 5.2}[len(paths)]
        for cell, f in zip(t.rows[0].cells, paths):
            _, n, ten, nguon, _ = FIG[f]
            w, h = Image.open(os.path.join(ROOT, f)).size
            wcm = min(width, 6.5 * w / h) if h > w else width
            if len(paths) == 1 and w > 2 * h:
                wcm = 16.0
            p = cell.paragraphs[0]; p.alignment = C
            p.add_run().add_picture(os.path.join(ROOT, f), width=Cm(wcm))
            p = cell.add_paragraph(); p.alignment = C
            self.runs(p, 'Hình %d' % n, 'b i', 10); self.runs(p, ': ' + ten + '.', 'i', 10)
            p = cell.add_paragraph(); p.alignment = C
            self.runs(p, 'Nguồn: ' + nguon, 'i gray', 9)

    def save(self, name):
        out = os.path.join(HERE, name); self.d.save(out); return out


def build(gv):
    D = Doc(gv)
    D.para([(S.HEAD, 'b blue')], C, 16)
    D.para([(S.LESSON, 'b red')], C, 13)
    if gv:
        D.para([('Bản GV [SOURCE]: đáp án và giải thích ngay sau từng câu hỏi', 'i gray')], C, 10)

    D.para([('1. Ôn tập nhanh', 'b blue')], L, 13)
    for t in S.INTRO:
        D.para([(t, '')])

    D.para([(S.PX_HEAD, 'b blue')], L, 13)
    for px in S.PX:
        code, name = px['title'].split('. ', 1)
        p = D.para([(code + '. ', 'b red'), (name, 'b red')], L, 13); D.border_bottom(p)
        D.para([('Lý thuyết: ', 'b'), (px['theory'], '')])
        for v in px['vd']:
            D.para([('VD: ', 'b'), (v, '')])
        figs = px.get('figs', [])
        for row in (figs if figs and isinstance(figs[0], list) else [figs] if figs else []):
            D.figures(row)
        for rf in px['reflex']:
            parts = re.split(r'→|thì nghĩ ngay:', rf, maxsplit=1)
            if len(parts) == 2:
                D.para([('⇒ ', 'b blue'), (parts[0].strip() + ' →', 'b'), (' ' + parts[1].strip(), '')], indent=0.4)
            else:
                D.para([('⇒ ', 'b blue'), (rf, '')], indent=0.4)
        traps = px['trap']
        if len(traps) == 1:
            D.para([('Bẫy: ', 'b org'), (traps[0], '')])
        else:
            D.para([('Bẫy:', 'b org')])
            for i, t in enumerate(traps):
                D.para([('(%s) ' % 'abcd'[i], 'b org'), (t, '')], indent=0.4)
        D.para([('Ngữ liệu lạ', 'b i')], L)
        for b in px['passage']:
            if isinstance(b, str):
                D.para([(b, '')])
            elif 'fig' in b:
                D.figures([b['fig']])
            else:
                D.table(b['rows'], b['label'], b['cap'])
        D.para([('Câu hỏi:', 'b')], L)
        for i, q in enumerate(px['questions'], 1):
            D.para([('(%d) ' % i, 'b org'), (q['q'], '')])
            if not gv:
                for _ in range(S.so_dong(q)):
                    D.d.add_paragraph()
            if gv:
                D.para([('▸ Đáp án: ', 'b grn'), (q['ans'], 'grn')])
                D.para([('▸ Giải thích: ', 'b grn'), (q['exp'], 'grn')])

    if gv:
        D.para([('CÂU HỎI TRANH BIỆN CỦA BUỔI', 'b red')], C, 13)
        for i, t in enumerate(S.DEBATE, 1):
            D.para([('%d. ' % i, 'b'), (t['q'], '')])
            D.para([('Gợi ý GV: ', 'b grn'), (t['hint'], 'grn')])
        D.table(
            [['PX', 'Phản xạ', 'Số câu']] +
            [[p['title'].split('. ', 1)[0], p['title'].split('. ', 1)[1], str(len(p['questions']))] for p in S.PX],
            '2.5', 'Tổng hợp phản xạ Buổi 2.')
        D.para([(S.SUMMARY, '')], L)
    D.para([('Nguồn: Nguyễn Nghĩa Đô biên soạn', 'b')], R)
    return D.save(('Buổi 2. [SOURCE] ' if gv else 'Buổi 2. ') + 'Le Chatelier qua Q-Kc.docx')


if __name__ == '__main__':
    print(build(False)); print(build(True))
